"""Builds the Anomaly Benchmark v2 documents (.txt and .docx) from the text below.

Every document is fictional and marked as synthetic. This file holds ONLY the
document text; the answer key lives in ground_truth.json and is never read by
the system under test.

    python build_fixtures.py        (needs python-docx)
"""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

import sys

HERE = Path(__file__).parent
SET = sys.argv[1] if len(sys.argv) > 1 else "main"   # "main" or "heldout"
OUT = (HERE / "documents") if SET == "main" else (HERE / "heldout" / "documents")
OUT.mkdir(parents=True, exist_ok=True)

BANNER = "SYNTHETIC TEST DOCUMENT - NOT FOR LEGAL USE. All names, addresses, amounts and identifiers are invented."

DOCS = {}

# --------------------------------------------------------------------------- B01 rental, flawed
DOCS["B01_residential_rental"] = """# RESIDENTIAL RENTAL AGREEMENT
{banner}

This Rental Agreement is made on 12 March 2027 between Mr. Sanjay Pillai of 3 Temple Street, Thrissur, Kerala (the "Landlord") and Ms. Meera Krishnan of 18 Garden Colony, Kochi, Kerala (the "Tenant").

1. PREMISES
The Landlord lets to the Tenant the second-floor apartment at 22 Canal Road, Kochi, Kerala 682016, for residential use by the Tenant and her immediate family only.

2. TERM
The tenancy runs for 11 months, from 1 April 2027 to 29 February 2028.

3. RENT
The Tenant shall pay monthly rent of INR 28,500, due on or before the 7th day of each month by bank transfer.

4. SECURITY DEPOSIT
The Tenant shall pay an interest-free security deposit of INR 57,000, equal to two months' rent, when signing this Agreement. The deposit will be refunded to the Tenant after the Tenant vacates the premises and returns the keys, less any proven damage beyond ordinary wear and tear.

5. UTILITIES
The Tenant pays electricity and water charges as billed. Society charges are paid by the Landlord.

6. REPAIRS
The Tenant shall bear minor repairs costing up to INR 2,000 each. All other repairs are the Landlord's responsibility.

7. TERMINATION
The Landlord may terminate this Agreement by giving the Tenant 15 days written notice. The Tenant may terminate this Agreement by giving the Landlord 3 months written notice.

8. ACCESS
The Landlord may enter the premises with reasonable prior notice to inspect or repair them.

9. JURISDICTION
The courts at Kochi, Kerala shall have jurisdiction over any dispute arising from this Agreement.

SCHEDULE OF PAYMENTS
Monthly Rent: INR 31,000
Security Deposit: INR 57,000
Maintenance charges: INR 1,500 per month
"""

# --------------------------------------------------------------------------- B02 leave and licence, flawed (single)
DOCS["B02_leave_and_license"] = """# LEAVE AND LICENSE AGREEMENT

Executed at Nagpur on 2 February 2027.

The Licensor, Mrs. Sunita Wagh, residing at 41 Civil Lines, Nagpur, Maharashtra (the "Licensor"), grants to Mr. Imran Qureshi, residing at 9 Dharampeth Extension, Nagpur, Maharashtra (the "Licensee"), a licence on the terms below.

Clause 1: Licensed Premises
Flat 5, Shanti Niketan, Ramdaspeth, Nagpur 440010, for the Licensee's own residential use.

Clause 2: Licence Period
The licence period begins on 1 April 2027 and ends on 31 March 2026. The licence may be renewed only by a signed writing.

Clause 3: Licence Fee
The Licensee shall pay a monthly licence fee of INR 18,000 on or before the 5th of every month.

Clause 4: Security Deposit
The Licensee shall pay a refundable security deposit of INR 54,000. The Licensor shall return it within 21 days after the Licensee hands back the keys, less documented dues.

Clause 5: Nature of Licence
This is a licence only. It creates no tenancy and gives the Licensee no interest in the premises.

Clause 6: Termination
Either party may end the licence by giving 45 days written notice to the other, after the first three months.

Clause 7: Governing Courts
The courts at Nagpur have exclusive jurisdiction.

{banner}
"""

# --------------------------------------------------------------------------- B03 land lease, flawed (multi)
DOCS["B03_land_lease"] = """# LEASE DEED (AGRICULTURAL LAND)
{banner}

This Deed of Lease is made on 20 January 2027 between Shri Ramesh Gaikwad of Village Pimpalgaon, Taluka Baramati, District Pune, Maharashtra (the "Lessor") and Shri Dinesh Jagtap of Village Morgaon, District Pune, Maharashtra (the "Lessee").

ARTICLE I - THE LAND
The Lessor leases to the Lessee the agricultural land at Survey No. 118/3, Village Pimpalgaon, Taluka Baramati, admeasuring 3.2 acres, with the well and pump-house standing on it.

ARTICLE II - TERM
The lease is for a term of 7 years, beginning on 1 June 2027 and ending on 31 May 2032.

ARTICLE III - PREMIUM AND RENT
The Lessee shall pay a one-time premium of INR 3,00,000 on signing, and an annual rent of INR 2,10,000 payable on 1 June each year.

ARTICLE IV - PERMITTED USE
The Lessee shall use the land only for cultivation of seasonal crops and shall not construct any permanent building on it.

ARTICLE V - BOUNDARIES
The land at Survey No. 118/5 admeasuring 3.8 acres is bounded on the east by the village road, on the west by the canal and on the north by the Lessor's remaining land.

ARTICLE VI - TERMINATION
Either party may terminate this Deed on 90 days written notice if the other party fails to cure a breach within 30 days of written notice.

ARTICLE VII - DISPUTES
Disputes shall first go to mediation at the Baramati taluka level and then to the civil courts at Pune.
"""

# --------------------------------------------------------------------------- B04 agreement to sell, flawed (multi)
DOCS["B04_agreement_to_sell"] = """# AGREEMENT TO SELL
{banner}

Dated 1 April 2027.

Between Mr. Anil Bhosale of 12 Kalyani Nagar, Pune (the "Seller") and Mrs. Rekha Iyer of 7 Aundh Road, Pune (the "Buyer").

1. Subject Property
The Seller agrees to sell Flat 14C on the fourteenth floor of Orion Residency, Baner, Pune 411045, with a carpet area of 920 sq. ft.

2. Price
The agreed price for the property is INR 64,00,000.

3. Payments
The Buyer has paid an advance of INR 8,00,000 on signing. The balance of INR 52,00,000 shall be paid by the Buyer on or before 30 June 2027 against registration of the sale deed.

4. Possession
The Seller shall hand over vacant possession of Flat 14D on 15 March 2027.

5. Title
The Seller states that the property is free from mortgage, lien or litigation and shall indemnify the Buyer against any undisclosed encumbrance.

6. Default
If the Buyer fails to pay the balance on time, the Seller may forfeit the advance after 15 days written notice.

7. Costs
The Buyer shall bear stamp duty and registration charges.

8. Courts
The courts at Pune have exclusive jurisdiction.
"""

# --------------------------------------------------------------------------- B05 employment contract, flawed (multi)
DOCS["B05_employment_contract"] = """# CONTRACT OF EMPLOYMENT

This contract is made on 10 May 2027 between Kestrel Data Systems Pvt. Ltd., 17 Outer Ring Road, Bengaluru (the "Company") and Mr. Farhan Siddiqui of 6 Koramangala 4th Block, Bengaluru (the "Employee").

1) DESIGNATION AND JOINING
The Company employs the Employee as Product Analyst. The Employee shall join on 1 July 2027.

2) REMUNERATION
The annual cost to the Company (CTC) is INR 1,800,000. It is made up of a basic salary of INR 60,000 per month, house rent allowance of INR 30,000 per month and a special allowance of INR 40,000 per month.

3) BONUS
The Employee is eligible for an annual performance bonus of up to 10% of CTC, paid after the annual review. Any bonus is at the Company's discretion.

4) PROBATION
The Employee shall serve a probation of 3 months from joining. Confirmation will be communicated in writing.

5) NOTICE
After confirmation, either party shall give the other 60 days written notice to end this contract.

6) CONFIDENTIALITY
The Employee shall keep the Company's non-public information confidential during and after employment.

7) INTELLECTUAL PROPERTY
Work produced by the Employee in the course of employment belongs to the Company.

8) TERMS OF SERVICE
The Employee's probation period shall be 6 months, and the performance bonus is 15% of CTC. The Employee is entitled to 20 days of paid leave per year.

9) JURISDICTION
Courts at Bengaluru have jurisdiction.

{banner}
"""

# --------------------------------------------------------------------------- B06 NDA, flawed (multi)
DOCS["B06_nda"] = """# CONFIDENTIALITY AGREEMENT
{banner}

This Agreement is made on 4 June 2027 between Helio Grid Energy Pvt. Ltd., 8 Sector 18, Noida (the "Disclosing Party") and Mr. Tarun Bedi, an independent consultant at 22 Lajpat Nagar, New Delhi (the "Recipient").

1. Purpose
The Disclosing Party will share confidential information with the Recipient only to evaluate a proposed grid-analytics engagement.

2. Definition
"Confidential Information" means non-public technical or business information disclosed by the Disclosing Party in any form.

3. Obligations
The Recipient shall keep the Confidential Information secret, use it only for the Purpose, and not disclose it to anyone else without written consent.

4. Exclusions
The obligations do not cover information that is public, already known to the Recipient, or independently developed.

5. Duration
The Recipient's confidentiality obligations continue for 3 years from the date of this Agreement.

6. Return of Information
Within 10 days of a written request, the Recipient shall return or destroy all Confidential Information. The Recipient may nevertheless keep archival copies for 6 years after the date of this Agreement.

7. Further Term
For the avoidance of doubt, the Recipient's duty of confidentiality continues for a period of 6 years from the date of this Agreement.

8. Jurisdiction
The courts at New Delhi have exclusive jurisdiction.
"""

# --------------------------------------------------------------------------- B07 service agreement, flawed (multi)
DOCS["B07_service_agreement"] = """# SERVICES AGREEMENT
{banner}

This Services Agreement is made on 18 July 2027 between Ironbridge Foods Pvt. Ltd., 31 Ring Road, Surat, Gujarat (the "Customer") and Pixelcraft Digital LLP, 5 Pal Gam Road, Surat, Gujarat (the "Supplier").

ARTICLE I - SERVICES
The Supplier shall design and operate an inventory dashboard for the Customer as described in the attached statement of work.

ARTICLE II - FEES
The Customer shall pay the Supplier a fixed fee of INR 9,00,000 for the build and INR 40,000 per month for support.

ARTICLE III - PAYMENT
The Supplier shall invoice monthly. The Customer shall pay each invoice within 30 days of receiving it.

ARTICLE IV - LATE INVOICES
Each invoice is payable within 45 days of the invoice date, and the Supplier may not charge interest before that date.

ARTICLE V - CONFIDENTIALITY
Each party shall protect the other's non-public information and use it only for this Agreement.

ARTICLE VI - LIABILITY AND INDEMNITY
The Supplier's total liability for claims arising from the Services is limited to the fees paid in the preceding 12 months. The Supplier shall indemnify the Customer without any limit against all losses arising from data breaches caused by the Supplier.

ARTICLE VII - TERMINATION
Either party may terminate this Agreement on 60 days written notice.

ARTICLE VIII - JURISDICTION
The courts at Surat have exclusive jurisdiction.
"""

# --------------------------------------------------------------------------- C01 sale deed, clean
DOCS["C01_sale_deed"] = """# SALE DEED
{banner}

This Sale Deed is executed on 14 August 2027 between Mr. Prakash Nair of 19 Indira Nagar, Thiruvananthapuram (the "Vendor") and Dr. Leena Thomas of 4 Kowdiar Road, Thiruvananthapuram (the "Purchaser").

ARTICLE I - THE PROPERTY
The Vendor sells to the Purchaser Flat 7A on the seventh floor of Lakeview Heights, Kowdiar, Thiruvananthapuram 695003, having a built-up area of 1,080 sq. ft., together with one covered parking space.

ARTICLE II - CONSIDERATION
The total consideration is INR 72,00,000. The Vendor acknowledges receipt of INR 9,00,000 earlier as advance and INR 63,00,000 today, making INR 72,00,000 in full.

ARTICLE III - TRANSFER OF TITLE
On execution of this Deed the Vendor transfers absolute ownership of Flat 7A to the Purchaser.

ARTICLE IV - POSSESSION
The Vendor hands over vacant possession of Flat 7A to the Purchaser on the date of this Deed.

ARTICLE V - ENCUMBRANCES
The Vendor declares that Flat 7A is free from any mortgage, charge or litigation and shall indemnify the Purchaser against any defect in title.

ARTICLE VI - COSTS
The Purchaser has paid the stamp duty and registration charges.

ARTICLE VII - COURTS
The courts at Thiruvananthapuram have jurisdiction over any dispute relating to this Deed.
"""

# --------------------------------------------------------------------------- C02 builder-buyer, clean
DOCS["C02_builder_buyer"] = """# AGREEMENT FOR ALLOTMENT
{banner}

Made on 5 September 2027 between Greenarc Realty Pvt. Ltd., 9 Sector 62, Noida (the "Promoter") and Mr. Vivek Arora of 15 Vasant Kunj, New Delhi (the "Allottee").

1. The Project
The Promoter is developing the project "Greenarc Square" at Sector 150, Noida, registered with the state real estate authority.

2. The Unit
The Promoter allots Unit 904 on the ninth floor of Tower C, with a carpet area of 640 sq. ft., to the Allottee.

3. Price
The total price of the unit is INR 85,00,000.

4. Payment Plan
The Allottee shall pay 10% on booking, 40% on completion of the foundation, 40% on completion of the superstructure and 10% on offer of possession.

5. Possession
The Promoter shall offer possession of Unit 904 by 31 March 2029.

6. Delay
If possession is delayed beyond the date above, the Promoter shall pay the Allottee interest at 9% per annum on the amounts paid for the period of delay.

7. Defects
The Promoter shall repair any structural defect notified within 5 years of possession, free of cost.

8. Cancellation
If the Allottee cancels, the Promoter shall refund the amounts paid, less a charge of 2% of the price, within 60 days.

9. Courts
The courts at Gautam Buddha Nagar have jurisdiction.
"""

# --------------------------------------------------------------------------- C03 offer letter, clean
DOCS["C03_offer_letter"] = """# OFFER OF EMPLOYMENT

Date: 25 May 2027

Dear Ms. Sneha Rao,

Brightline Logistics Pvt. Ltd., 70 Anna Salai, Chennai (the "Company") is pleased to offer you the position of Operations Associate, reporting to the Operations Manager.

Your joining date will be 14 June 2027.

Your annual cost to the Company (CTC) is INR 6,00,000, made up of a basic salary of INR 25,000 per month, house rent allowance of INR 15,000 per month and special allowance of INR 10,000 per month.

You will be on probation for 3 months from your joining date. Either party may end your employment during probation on 15 days written notice.

Please sign and return a copy of this letter by 5 June 2027 to accept this offer.

{banner}
"""

# --------------------------------------------------------------------------- C04 appointment letter, clean
DOCS["C04_appointment_letter"] = """{banner}

# LETTER OF APPOINTMENT

Issued on 1 June 2027 by Meridian Public School Society, 12 Civil Lines, Jaipur (the "Employer") to Mr. Gaurav Saxena of 33 Malviya Nagar, Jaipur (the "Employee").

1. APPOINTMENT
You are appointed as Senior Teacher (Mathematics), reporting to the Principal.

2. JOINING
You shall join on 1 July 2027.

3. SALARY
Your gross monthly salary is INR 62,000, payable on the last working day of each month.

4. PROBATION
You will be on probation for 6 months from the date of joining.

5. NOTICE
After confirmation, either party may end this appointment by giving 60 days written notice or salary in lieu of notice.

6. CONFIDENTIALITY
You shall keep confidential all examination material and student records you handle.
"""

# --------------------------------------------------------------------------- C05 mutual NDA, clean, "Party A"/"Party B"
DOCS["C05_mutual_nda"] = """# MUTUAL NON-DISCLOSURE AGREEMENT
{banner}

This Agreement is made on 9 August 2027 between Tessera Biosystems Pvt. Ltd., 14 Genome Valley, Hyderabad ("Party A") and Calder Instruments Pvt. Ltd., 60 Peenya Industrial Area, Bengaluru ("Party B").

1. Purpose
Each party may disclose confidential information to the other solely to assess a joint instrument development project.

2. Confidential Information
"Confidential Information" means non-public information disclosed by either party in any form that is marked confidential or would reasonably be understood to be confidential.

3. Obligations
Each party shall keep the other party's Confidential Information secret, use it only for the Purpose, and disclose it only to employees who need it and are bound by equivalent duties.

4. Exclusions
These obligations do not apply to information that is public, was lawfully known earlier, or is independently developed.

5. Term
This Agreement remains in force for 2 years from the date above. Each party's confidentiality duties continue for 3 years after it ends.

6. Return
Within 15 days of a written request, each party shall return or destroy the other's Confidential Information.

7. Jurisdiction
The courts at Bengaluru have exclusive jurisdiction.
"""

# --------------------------------------------------------------------------- C06 MSA with explicit cap and carve-outs, clean
DOCS["C06_msa"] = """# MASTER SERVICES AGREEMENT
{banner}

This Master Services Agreement is made on 23 August 2027 between Coastal Freight Lines Pvt. Ltd., 3 Port Road, Visakhapatnam (the "Customer") and Nimbus Cloud Services Pvt. Ltd., 21 Madhapur, Hyderabad (the "Supplier").

1. Services
The Supplier shall provide cloud hosting and support services described in each statement of work signed under this Agreement.

2. Fees
The Customer shall pay the fees in each statement of work. The first statement of work fixes a monthly fee of INR 2,50,000.

3. Payment
The Supplier shall invoice monthly. The Customer shall pay each undisputed invoice within 30 days of receipt.

4. Confidentiality
Each party shall keep the other's non-public information confidential during the term and for 3 years afterwards.

5. Limitation of Liability
Each party's total liability under this Agreement is limited to the fees paid in the 12 months before the claim. This limit does not apply to liability for breach of confidentiality or for wilful misconduct.

6. Term and Termination
This Agreement continues for 3 years from the date above. Either party may terminate it on 90 days written notice.

7. Disputes
The parties shall escalate any dispute to senior management and, failing resolution in 30 days, refer it to arbitration in Hyderabad.

8. Jurisdiction
The courts at Hyderabad have exclusive jurisdiction.
"""

# --------------------------------------------------------------------------- C07 rental, clean, inline-heading style
DOCS["C07_residential_rental"] = """# RENT AGREEMENT
{banner}

Made on 6 October 2027 between Mrs. Latha Venkatesh of 27 Adyar Main Road, Chennai (the "Landlord") and Mr. Ajay Menon of 5 Velachery Road, Chennai (the "Tenant").

1. Premises: The Landlord lets the first-floor flat at 12 Besant Avenue, Chennai 600090 to the Tenant for residential use.

2. Term: The tenancy runs for 11 months from 1 November 2027 to 30 September 2028.

3. Rent: The Tenant shall pay rent of INR 22,000 per month on or before the 5th of each month.

4. Security Deposit: The Tenant shall pay a refundable deposit of INR 66,000, equal to three months' rent. The Landlord shall refund it within 30 days after the Tenant vacates and returns the keys, less proven damage.

5. Maintenance: The Tenant pays electricity and water charges. The Landlord pays property tax and structural repairs.

6. Termination: Either party may end this tenancy by giving the other 2 months written notice.

7. Courts: The courts at Chennai have jurisdiction.
"""

if SET == "heldout":
    DOCS = {}
    exec((HERE / "heldout" / "heldout_texts.py").read_text(encoding="utf-8"))

for name, body in DOCS.items():
    text = body.format(banner=BANNER).strip() + "\n"
    (OUT / f"{name}.txt").write_text(text, encoding="utf-8")
    d = Document()
    d.styles["Normal"].font.name = "Times New Roman"
    d.styles["Normal"].font.size = Pt(11)
    for line in text.splitlines():
        if not line.strip():
            continue
        if line.startswith("# "):
            p = d.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(line[2:])
            r.bold = True
            r.font.size = Pt(14)
        else:
            d.add_paragraph(line)
    d.save(OUT / f"{name}.docx")
    print(name)
