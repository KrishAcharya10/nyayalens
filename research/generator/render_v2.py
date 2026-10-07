"""
Renders v2 documents to .pdf (native text or simulated scan) and .docx.

Fixes a v1 defect: render.py used reportlab's built-in Helvetica, which
has no rupee glyph, so every "₹" in v1's PDFs came out of text extraction
mangled (the root cause of the currency workaround in the app's rule
patterns and services/ml/metrics.py). v2 embeds a TrueType font that has
U+20B9 and refuses to fall back to one that doesn't.

Layout also varies per document (font, size, alignment, margins), and a
configurable fraction of PDFs are rendered as page images with scan-like
degradation, so the OCR path gets exercised by data with ground truth.
"""

import io
import os
import platform
from dataclasses import dataclass

from docx import Document as DocxDocument
from docx.shared import Pt
from PIL import Image, ImageDraw, ImageFilter, ImageFont
from reportlab.lib.enums import TA_JUSTIFY, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

from value_pools_v2 import rng

RUPEE = 0x20B9

# (family name, regular file, bold file, docx font name)
_FONT_CANDIDATES = {
    "Windows": [
        ("Times", "times.ttf", "timesbd.ttf", "Times New Roman"),
        ("Arial", "arial.ttf", "arialbd.ttf", "Arial"),
        ("Georgia", "georgia.ttf", "georgiab.ttf", "Georgia"),
        ("Calibri", "calibri.ttf", "calibrib.ttf", "Calibri"),
    ],
    "Linux": [
        ("DejaVuSans", "DejaVuSans.ttf", "DejaVuSans-Bold.ttf", "DejaVu Sans"),
        ("DejaVuSerif", "DejaVuSerif.ttf", "DejaVuSerif-Bold.ttf", "DejaVu Serif"),
    ],
}
_FONT_DIRS = {
    "Windows": [os.path.join(os.environ.get("WINDIR", r"C:\Windows"), "Fonts")],
    "Linux": ["/usr/share/fonts/truetype/dejavu", "/usr/share/fonts/dejavu"],
}


@dataclass
class FontChoice:
    family: str
    regular_path: str
    bold_family: str
    docx_name: str


_fonts: list[FontChoice] | None = None


def available_fonts() -> list[FontChoice]:
    """Registers every candidate font that exists and has the rupee glyph."""
    global _fonts
    if _fonts is not None:
        return _fonts
    _fonts = []
    system = platform.system()
    for family, regular, bold, docx_name in _FONT_CANDIDATES.get(system, []):
        for d in _FONT_DIRS.get(system, []):
            reg_path, bold_path = os.path.join(d, regular), os.path.join(d, bold)
            if not os.path.exists(reg_path):
                continue
            font = TTFont(family, reg_path)
            if RUPEE not in font.face.charToGlyph:
                continue
            pdfmetrics.registerFont(font)
            bold_family = family
            if os.path.exists(bold_path):
                pdfmetrics.registerFont(TTFont(f"{family}-Bold", bold_path))
                bold_family = f"{family}-Bold"
            _fonts.append(FontChoice(family, reg_path, bold_family, docx_name))
            break
    if not _fonts:
        raise RuntimeError(
            "No TrueType font with the rupee glyph (U+20B9) was found (looked for Times/Arial/"
            "Georgia/Calibri on Windows, DejaVu on Linux) - refusing to render PDFs that would "
            "silently mangle every currency amount, which is the v1 defect this renderer fixes."
        )
    return _fonts


@dataclass
class Layout:
    font: FontChoice
    size: float
    justify: bool
    margin_in: float


def random_layout() -> Layout:
    return Layout(
        font=rng.choice(available_fonts()),
        size=rng.choice([10, 10.5, 11, 11.5, 12]),
        justify=rng.random() < 0.4,
        margin_in=rng.choice([0.8, 1.0, 1.2]),
    )


def _is_heading(line: str) -> bool:
    s = line.strip()
    return 0 < len(s) <= 70 and s.upper() == s and any(c.isalpha() for c in s)


def render_pdf(text: str, out_path: str, layout: Layout) -> None:
    body = ParagraphStyle("body", fontName=layout.font.family, fontSize=layout.size,
                          leading=layout.size * 1.35, alignment=TA_JUSTIFY if layout.justify else TA_LEFT)
    head = ParagraphStyle("head", parent=body, fontName=layout.font.bold_family, alignment=TA_LEFT)
    margin = layout.margin_in * inch
    doc = SimpleDocTemplate(out_path, pagesize=A4, leftMargin=margin, rightMargin=margin,
                            topMargin=margin, bottomMargin=margin)
    story = []
    for line in text.split("\n"):
        if not line.strip():
            story.append(Spacer(1, layout.size * 0.8))
            continue
        safe = line.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        story.append(Paragraph(safe, head if _is_heading(line) else body))
    doc.build(story)


def render_docx(text: str, out_path: str, layout: Layout) -> None:
    doc = DocxDocument()
    style = doc.styles["Normal"]
    style.font.name = layout.font.docx_name
    style.font.size = Pt(layout.size)
    for line in text.split("\n"):
        para = doc.add_paragraph()
        run = para.add_run(line)
        run.bold = _is_heading(line)
    doc.save(out_path)


# --- simulated scans ----------------------------------------------------------

_DPI = 150
_PAGE_PX = (int(8.27 * _DPI), int(11.69 * _DPI))


def _wrap(draw: ImageDraw.ImageDraw, line: str, font, width: int) -> list[str]:
    words, out, cur = line.split(" "), [], ""
    for w in words:
        trial = f"{cur} {w}".strip()
        if draw.textlength(trial, font=font) <= width or not cur:
            cur = trial
        else:
            out.append(cur)
            cur = w
    out.append(cur)
    return out


def _degrade(page: Image.Image) -> Image.Image:
    page = page.rotate(rng.uniform(-1.2, 1.2), resample=Image.BICUBIC, fillcolor=255)
    page = page.filter(ImageFilter.GaussianBlur(rng.uniform(0.2, 0.9)))
    px = page.load()
    w, h = page.size
    for _ in range(rng.randint(300, 1500)):  # speckle
        px[rng.randrange(w), rng.randrange(h)] = rng.randint(0, 120)
    buf = io.BytesIO()
    page.save(buf, format="JPEG", quality=rng.randint(35, 75))  # compression artefacts
    buf.seek(0)
    return Image.open(buf).convert("L")


def render_scanned_pdf(text: str, out_path: str, layout: Layout) -> None:
    """Image-only PDF: no text layer, so the app must OCR it."""
    font_px = int(layout.size * _DPI / 72)
    font = ImageFont.truetype(layout.font.regular_path, font_px)
    margin = int(layout.margin_in * _DPI)
    width = _PAGE_PX[0] - 2 * margin
    line_h = int(font_px * 1.4)
    paper = rng.randint(232, 252)

    pages, page, draw, y = [], None, None, margin
    for line in text.split("\n"):
        for piece in (_wrap(ImageDraw.Draw(Image.new("L", (1, 1))), line, font, width) if line else [""]):
            if page is None or y + line_h > _PAGE_PX[1] - margin:
                if page is not None:
                    pages.append(_degrade(page))
                page = Image.new("L", _PAGE_PX, paper)
                draw = ImageDraw.Draw(page)
                y = margin
            if piece:
                draw.text((margin, y), piece, font=font, fill=rng.randint(0, 40))
            y += line_h
    if page is not None:
        pages.append(_degrade(page))
    pages[0].save(out_path, "PDF", resolution=_DPI, save_all=True, append_images=pages[1:])
